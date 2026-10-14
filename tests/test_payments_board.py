"""Экран «Оплаты»: кто оплатил за месяц, ручная отметка ГП, отмена, возврат.

`services/payment_board.py`, `payments.mark_paid_manually`,
`payments.revert_paid`, роуты `api/cabinet_payments_admin.py`. План —
`plans/2026-09-30-apparchi-monthly-payments.md`, «Ручная отметка оплаты» и «Возврат».
"""
from datetime import date, datetime, timezone

import pytest

from app.models.audit_log import AuditLog
from app.models.notification import Notification
from app.models.payment import (
    COHORT_FROM,
    SOURCE_MANUAL,
    SOURCE_PRODAMUS,
    STATUS_CANCELLED,
    STATUS_PAID,
    STATUS_PENDING,
    STATUS_REFUNDED,
    Payment,
    PaymentPrice,
)
from app.services import payment_board as pb
from app.services import payments as ps
from app.services.tz import MSK_TZ

OCT = date(2026, 10, 1)
SEP = date(2026, 9, 1)
NOV = date(2026, 11, 1)


def _msk(*args) -> datetime:
    return datetime(*args, tzinfo=MSK_TZ).astimezone(timezone.utc)


@pytest.fixture()
def prices(db):
    db.add(PaymentPrice(tariff="Я С ВАМИ", cohort=COHORT_FROM, amount_kop=1325500))
    db.commit()


@pytest.fixture()
def people(db, user_factory):
    chief = user_factory(vk_id=990_001, name="Главный", is_admin=True, role_name="суперадмин")
    curator = user_factory(vk_id=990_002, name="Куратор", role_name="куратор")
    return {"chief": chief, "curator": curator}


@pytest.fixture()
def make_payer(db, user_factory):
    counter = iter(range(990_100, 990_999))

    def _make(name="Иванова Анна", window=(10, 15), **fields):
        student = user_factory(vk_id=next(counter), name=name, tariff="Я С ВАМИ")
        student.pay_window_start, student.pay_window_end = window
        student.pay_cohort = COHORT_FROM
        for key, value in fields.items():
            setattr(student, key, value)
        db.commit()
        return student
    return _make


@pytest.fixture()
def no_delivery(monkeypatch):
    async def _noop(ids):
        return None
    monkeypatch.setattr("app.api.cabinet_payments_admin.notify_many", _noop)


def _mark(db, student, by, period=OCT, amount=1325500):
    payment, notes = ps.mark_paid_manually(
        db, student, period=period, amount_kop=amount, paid_on=date(2026, 10, 12),
        comment="платёжка 15", performed_by_id=by.id,
    )
    db.commit()
    return payment, notes


# ── Кто оплатил ──────────────────────────────────────────────────────────────

@pytest.mark.parametrize("now, state, detail", [
    (_msk(2026, 10, 5, 12, 0), pb.STATE_UPCOMING, "Окно откроется 10.10"),
    (_msk(2026, 10, 12, 12, 0), pb.STATE_OPEN, "Окно до 15.10"),
    (_msk(2026, 10, 16, 10, 0), pb.STATE_OVERDUE, "Окно закрылось 16.10"),
])
def test_unpaid_state_follows_window(db, prices, make_payer, now, state, detail):
    make_payer()

    row = pb.month_board(db, OCT, now).rows[0]

    assert (row.state, row.detail, row.amount_kop) == (state, detail, 1325500)


def test_february_window_ends_on_last_day(db, prices, make_payer):
    """Окно 25–30 в феврале идёт по 28-е и закрывается 1 марта в 2:00."""
    make_payer(window=(25, 30))
    feb = date(2027, 2, 1)
    assert pb.month_board(db, feb, _msk(2027, 2, 27, 12, 0)).rows[0].detail == "Окно до 28.02"
    assert pb.month_board(db, feb, _msk(2027, 3, 1, 1, 59)).rows[0].state == pb.STATE_OPEN
    assert pb.month_board(db, feb, _msk(2027, 3, 1, 2, 0)).rows[0].state == pb.STATE_OVERDUE


def test_month_paid_in_settings_counts_as_paid(db, prices, make_payer, people):
    student = make_payer()
    ps.apply_payment_settings(db, people["chief"].id, student, window_start=10, window_end=15,
                              cohort=COHORT_FROM, price_kop=None, paid_through=NOV)
    db.commit()

    board = pb.month_board(db, OCT, _msk(2026, 10, 20, 12, 0))

    assert (board.rows[0].state, board.rows[0].detail) == (pb.STATE_PAID, "Отмечено без платежа")
    assert board.counts[pb.STATE_PAID] == 1
    # Денег по такому месяцу нет: в итоги он идёт отдельным счётчиком
    # (сентябрь 2026 — техническая отметка запуска у всех 91).
    assert (board.paid_count, board.paid_total_kop, board.no_payment_count) == (0, 0, 1)
    assert board.rows[0].no_payment


# ── Статистика оплат (план 2026-10-13-payment-stats) ─────────────────────────

def _prodamus_paid(db, student, *, tariff="Я С ВАМИ", amount=1325500, period=OCT, order="777", note=None):
    payment = Payment(
        user_id=student.id, period=period, amount_kop=amount, paid_sum_kop=amount, tariff=tariff,
        status=STATUS_PAID, source=SOURCE_PRODAMUS, prodamus_order_id=order,
        paid_at=_msk(2026, 10, 13, 20, 25), note=note,
    )
    db.add(payment)
    db.commit()
    return payment


def test_paid_row_shows_tariff_of_payment(db, prices, make_payer):
    """Оплата сменила тариф («Тариф после оплаты»): строка месяца — с тарифом
    платежа, а не с нынешним тарифом ученика."""
    student = make_payer(email="anna@example.com", tg_username="@anna")
    _prodamus_paid(db, student, tariff="УВЕРЕННЫЙ МАКСИМУМ", amount=1875500,
                   note="тариф сменился с этой оплатой, был Я с вами")
    student.tariff = "Я САМ"
    db.commit()

    row = pb.month_board(db, OCT, _msk(2026, 10, 14, 12, 0)).rows[0]

    assert (row.tariff, row.amount_kop) == ("УВЕРЕННЫЙ МАКСИМУМ", 1875500)
    assert (row.paid_on, row.method) == ("13.10.2026", "Продамус")
    assert pb.contacts_text(student) == "anna@example.com · @anna"


def test_board_totals_by_tariff_and_debt(db, prices, make_payer, people):
    paid = make_payer("Аникина Ая")
    _prodamus_paid(db, paid)
    late = make_payer("Борисова Бэла", window=(1, 5))
    waiting = make_payer("Волкова Вера", window=(20, 25))
    marked = make_payer("Глебова Гуля")
    ps.apply_payment_settings(db, people["chief"].id, marked, window_start=10, window_end=15,
                              cohort=COHORT_FROM, price_kop=None, paid_through=OCT)
    other = make_payer("Дёмина Дарья", window=(1, 5), tariff="Я САМ", pay_price_kop=677500)
    db.commit()

    board = pb.month_board(db, OCT, _msk(2026, 10, 14, 12, 0))

    assert (board.paid_count, board.paid_total_kop, board.no_payment_count) == (1, 1325500, 1)
    assert (board.counts[pb.STATE_OVERDUE], board.debt_total_kop) == (2, 1325500 + 677500)
    totals = {t.tariff: t for t in board.by_tariff}
    assert (totals["Я С ВАМИ"].paid_count, totals["Я С ВАМИ"].paid_kop) == (1, 1325500)
    assert (totals["Я С ВАМИ"].debt_count, totals["Я С ВАМИ"].no_payment_count) == (1, 1)
    assert totals["Я С ВАМИ"].waiting_count == 1
    assert (totals["Я САМ"].debt_count, totals["Я САМ"].debt_kop) == (1, 677500)
    assert {late.id, other.id} == {r.student.id for r in board.rows if r.state == pb.STATE_OVERDUE}
    assert waiting.id in {r.student.id for r in board.rows if r.state == pb.STATE_UPCOMING}


def test_months_summary_from_first_payment_to_now(db, prices, make_payer, people):
    """Сводка — от первого месяца с платежом, свежий первым. Сентябрь,
    закрытый одной отметкой, в неё не попадает."""
    student = make_payer()
    ps.apply_payment_settings(db, people["chief"].id, student, window_start=10, window_end=15,
                              cohort=COHORT_FROM, price_kop=None, paid_through=SEP)
    db.commit()
    _prodamus_paid(db, student)

    summary = pb.months_summary(db, _msk(2026, 11, 20, 12, 0))

    assert [m.period for m in summary] == [NOV, OCT]
    assert (summary[1].paid_count, summary[1].paid_kop) == (1, 1325500)
    assert (summary[0].debt_count, summary[0].debt_kop) == (1, 1325500)


def test_months_summary_without_payments_is_current_month(db, prices, make_payer):
    make_payer()
    assert [m.period for m in pb.months_summary(db, _msk(2026, 10, 12, 12, 0))] == [OCT]


def test_month_xlsx_has_four_sheets(db, prices, make_payer):
    import io

    import openpyxl

    student = make_payer("Иванова Анна", email="anna@example.com", tg_username="anna")
    _prodamus_paid(db, student, order="12345")
    make_payer("Петрова Поля", window=(1, 5))
    now = _msk(2026, 10, 14, 12, 0)

    content = pb.build_month_xlsx(pb.month_board(db, OCT, now), pb.months_summary(db, now))

    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Ученики", "По тарифам", "Должники", "По месяцам"]
    students = [[c.value for c in row] for row in wb["Ученики"].iter_rows(min_row=2)]
    anna = next(r for r in students if r[0] == "Иванова Анна")
    assert anna[1:4] == ["anna@example.com", "@anna", pb.tariff_text("Я С ВАМИ")]
    assert "12345" in anna and 13255 in anna and "13.10.2026" in anna
    assert [r[0].value for r in wb["Должники"].iter_rows(min_row=2)] == ["Петрова Поля"]
    assert wb["По месяцам"]["A2"].value == "Октябрь 2026"


def test_board_lists_only_accounted_students(db, prices, make_payer, monkeypatch):
    paying = make_payer("Иванова Анна")
    no_window = make_payer("Петров Пётр", window=(None, None))
    make_payer("Архивная Вера", archived_at=datetime.now(timezone.utc))
    make_payer("Блок Борис", is_active=False)
    service = make_payer("Служебный Роман")
    monkeypatch.setattr(pb, "REPORT_EXCLUDED_USER_IDS", {service.id})

    board = pb.month_board(db, OCT, _msk(2026, 10, 12, 12, 0))

    assert [r.student.id for r in board.rows] == [paying.id]
    assert [s.id for s in board.no_window] == [no_window.id]


def test_rows_sorted_overdue_first(db, prices, make_payer, people):
    paid = make_payer("Аникина Ая")
    late = make_payer("Яковлева Яна", window=(1, 5))
    _mark(db, paid, people["chief"])

    rows = pb.month_board(db, OCT, _msk(2026, 10, 12, 12, 0)).rows

    assert [(r.student.id, r.state) for r in rows] == [(late.id, pb.STATE_OVERDUE), (paid.id, pb.STATE_PAID)]


def test_parse_month_defaults_to_current_moscow_month():
    # 31.10 в 23:30 по Москве — уже не сентябрь и ещё не ноябрь по UTC-ошибке.
    assert pb.parse_month("", _msk(2026, 10, 31, 23, 30)) == OCT
    assert pb.parse_month("мусор", _msk(2026, 10, 31, 23, 30)) == OCT
    assert pb.parse_month("2026-09", _msk(2026, 10, 1, 0, 0)) == SEP


# ── Ручная отметка ───────────────────────────────────────────────────────────

def test_manual_mark_pays_and_remembers_previous_term(db, prices, make_payer, people):
    student = make_payer()

    payment, notes = _mark(db, student, people["chief"])
    db.refresh(student)

    assert (payment.status, payment.source, payment.marked_by_id) == (STATUS_PAID, SOURCE_MANUAL, people["chief"].id)
    assert payment.paid_until_before is None
    assert ps.paid_month(student) == OCT
    assert "платёжка 15" in payment.note and "поступил 12.10.2026" in payment.note
    assert any(n.user_id == student.id for n in notes)


@pytest.mark.parametrize("window, message", [((None, None), "окно оплаты"), ((10, 15), "уже оплачен")])
def test_manual_mark_refused(db, prices, make_payer, people, window, message):
    student = make_payer(window=window)
    if window[0]:
        _mark(db, student, people["chief"])

    with pytest.raises(ps.PaymentSettingsError, match=message):
        _mark(db, student, people["chief"])


def test_manual_mark_cancels_pending_link(db, prices, make_payer, people):
    student = make_payer()
    pending = Payment(user_id=student.id, period=OCT, amount_kop=1325500, status=STATUS_PENDING)
    db.add(pending)
    db.commit()

    _mark(db, student, people["chief"])
    db.refresh(pending)

    assert pending.status == STATUS_CANCELLED


# ── Отмена и возврат ─────────────────────────────────────────────────────────

def test_cancel_restores_term_from_settings(db, prices, make_payer, people):
    student = make_payer()
    ps.apply_payment_settings(db, people["chief"].id, student, window_start=10, window_end=15,
                              cohort=COHORT_FROM, price_kop=None, paid_through=SEP)
    db.commit()
    payment, _ = _mark(db, student, people["chief"])

    ps.revert_paid(db, payment, refund=False, performed_by_id=people["chief"].id, now=datetime.now(timezone.utc))
    db.commit()
    db.refresh(student)

    assert payment.status == STATUS_CANCELLED
    assert ps.paid_month(student) == SEP
    assert db.query(AuditLog).filter(AuditLog.action == "payment_mark_cancelled").count() == 1


def test_cancel_without_previous_term_leaves_student_untracked(db, prices, make_payer, people):
    student = make_payer()
    payment, _ = _mark(db, student, people["chief"])

    ps.revert_paid(db, payment, refund=False, performed_by_id=people["chief"].id, now=datetime.now(timezone.utc))
    db.commit()
    db.refresh(student)

    assert student.paid_until is None


def test_refund_without_previous_term_leaves_previous_month_paid(db, prices, make_payer, people):
    """Иначе ученик выпал бы из оплаты, и кабинет по неоплате не закрылся бы."""
    student = make_payer()
    payment, _ = _mark(db, student, people["chief"])

    ps.revert_paid(db, payment, refund=True, performed_by_id=people["chief"].id, now=datetime.now(timezone.utc))
    db.commit()
    db.refresh(student)

    assert (payment.status, payment.refunded_by_id) == (STATUS_REFUNDED, people["chief"].id)
    assert ps.paid_month(student) == SEP


def test_revert_of_older_month_keeps_later_payment(db, prices, make_payer, people):
    student = make_payer()
    october, _ = _mark(db, student, people["chief"], period=OCT)
    _mark(db, student, people["chief"], period=NOV)

    ps.revert_paid(db, october, refund=True, performed_by_id=people["chief"].id, now=datetime.now(timezone.utc))
    db.commit()
    db.refresh(student)

    assert ps.paid_month(student) == NOV


def test_prodamus_payment_refund_only(db, prices, make_payer, people):
    student = make_payer()
    payment = Payment(user_id=student.id, period=OCT, amount_kop=1325500, status=STATUS_PENDING,
                      source=SOURCE_PRODAMUS)
    db.add(payment)
    db.flush()
    ps.record_paid(db, payment, now=datetime.now(timezone.utc), source=SOURCE_PRODAMUS,
                   paid_sum_kop=1325500, performed_by_id=student.id)
    db.commit()

    with pytest.raises(ps.PaymentSettingsError, match="только возвратом"):
        ps.revert_paid(db, payment, refund=False, performed_by_id=people["chief"].id, now=datetime.now(timezone.utc))
    ps.revert_paid(db, payment, refund=True, performed_by_id=people["chief"].id, now=datetime.now(timezone.utc))
    assert payment.status == STATUS_REFUNDED


# ── Роуты и права ────────────────────────────────────────────────────────────

def test_page_shows_month(db, client, session_factory, prices, make_payer, people):
    make_payer("Иванова Анна")
    client.cookies.set("session_id", session_factory(people["chief"]).id)

    resp = client.get("/cabinet/superadmin/payments?month=2026-10")

    assert resp.status_code == 200
    assert "Октябрь 2026" in resp.text and "Иванова Анна" in resp.text
    assert "Отметить оплату" in resp.text
    assert "По тарифам" in resp.text and "По месяцам" in resp.text
    assert "/cabinet/superadmin/payments/export?month=2026-10" in resp.text


def test_export_route_returns_xlsx(db, client, session_factory, prices, make_payer, people):
    make_payer("Иванова Анна")
    client.cookies.set("session_id", session_factory(people["chief"]).id)

    resp = client.get("/cabinet/superadmin/payments/export?month=2026-10")

    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "2026-10.xlsx" in resp.headers["content-disposition"]


def test_routes_mark_cancel_refund(db, client, session_factory, prices, make_payer, people, no_delivery):
    student = make_payer()
    client.cookies.set("session_id", session_factory(people["chief"]).id)

    marked = client.post("/cabinet/superadmin/payments/mark", data={
        "user_id": student.id, "month": "2026-10", "amount": "13 255", "paid_on": "2026-10-12", "comment": "",
    })
    assert marked.status_code == 200, marked.text
    payment_id = marked.json()["payment_id"]
    assert db.query(Notification).filter(Notification.user_id == student.id).count() == 1

    again = client.post("/cabinet/superadmin/payments/mark", data={
        "user_id": student.id, "month": "2026-10", "amount": "13255", "paid_on": "2026-10-12",
    })
    assert again.status_code == 400 and "уже оплачен" in again.json()["detail"]

    assert client.post(f"/cabinet/superadmin/payments/{payment_id}/cancel").status_code == 200
    assert client.post(f"/cabinet/superadmin/payments/{payment_id}/refund").status_code == 404


@pytest.mark.parametrize("form, message", [
    ({"month": "", "amount": "1", "paid_on": "2026-10-12"}, "месяц"),
    ({"month": "2026-10", "amount": "", "paid_on": "2026-10-12"}, "сумму"),
    ({"month": "2026-10", "amount": "1", "paid_on": ""}, "дату"),
])
def test_mark_route_refuses_bad_form(db, client, session_factory, prices, make_payer, people, form, message):
    student = make_payer()
    client.cookies.set("session_id", session_factory(people["chief"]).id)

    resp = client.post("/cabinet/superadmin/payments/mark", data={"user_id": student.id, **form})

    assert resp.status_code == 400 and message in resp.json()["detail"]


def test_curator_has_no_access(db, client, session_factory, prices, make_payer, people):
    student = make_payer()
    client.cookies.set("session_id", session_factory(people["curator"]).id)

    page = client.get("/cabinet/superadmin/payments", follow_redirects=False)
    export = client.get("/cabinet/superadmin/payments/export", follow_redirects=False)
    mark = client.post("/cabinet/superadmin/payments/mark", headers={"Accept": "application/json"}, data={
        "user_id": student.id, "month": "2026-10", "amount": "1", "paid_on": "2026-10-12",
    })

    assert page.status_code in (302, 303, 403)
    assert export.status_code in (302, 303, 403)
    assert mark.status_code == 403
    assert db.query(Payment).count() == 0


def test_mark_refused_for_archived(db, client, session_factory, prices, make_payer, people):
    student = make_payer(archived_at=datetime.now(timezone.utc))
    client.cookies.set("session_id", session_factory(people["chief"]).id)

    resp = client.post("/cabinet/superadmin/payments/mark", data={
        "user_id": student.id, "month": "2026-10", "amount": "1", "paid_on": "2026-10-12",
    })

    assert resp.status_code == 409


# ── ГП: всё про оплату, как у суперадмина (владелец 14.10.2026) ────────────

def test_head_teacher_has_everything_about_payments(db, client, session_factory, prices, make_payer, user_factory):
    """Владелец 14.10.2026: «только ГП открыть всё по оплатам, что есть у СА» —
    экран, Excel, «Оплата списком», группа в карточке, отметка."""
    head = user_factory(vk_id=990_050, name="ГП", is_admin=True, role_name="админ")
    student = make_payer("Иванова Анна")
    client.cookies.set("session_id", session_factory(head).id)

    users_page = client.get("/cabinet/superadmin/users")
    assert users_page.status_code == 200
    assert "/cabinet/superadmin/payments" in users_page.text
    assert "/cabinet/superadmin/payment-import" in users_page.text

    page = client.get("/cabinet/superadmin/payments?month=2026-10")
    assert page.status_code == 200 and "Иванова Анна" in page.text
    export = client.get("/cabinet/superadmin/payments/export?month=2026-10")
    assert export.status_code == 200 and export.content[:2] == b"PK"
    assert client.get("/cabinet/superadmin/payment-import").status_code == 200

    manage = client.get(f"/cabinet/students/{student.id}/profile").json()["student"]["manage"]
    assert manage["payment"] is not None

    json_headers = {"Accept": "application/json"}
    resp = client.post("/cabinet/superadmin/payments/mark", headers=json_headers, data={
        "user_id": student.id, "month": "2026-10", "amount": "1", "paid_on": "2026-10-12",
    })
    assert resp.status_code == 200, resp.text
    assert db.query(Payment).filter(Payment.user_id == student.id).count() == 1


@pytest.mark.parametrize("role_name", ["модератор", "куратор"])
def test_moderator_and_curator_see_no_payments(db, client, session_factory, prices, make_payer, user_factory, role_name):
    """Оплата — по роли: у модератора уровень ГП (`rbac.effective_role_rank`),
    но ни экрана, ни карточки, ни одного адреса оплаты ему не открыто."""
    staff = user_factory(vk_id=990_060, name="Сотрудник", role_name=role_name)
    student = make_payer("Иванова Анна")
    client.cookies.set("session_id", session_factory(staff).id)
    json_headers = {"Accept": "application/json"}

    assert client.get("/cabinet/superadmin/payments?month=2026-10", headers=json_headers).status_code == 403
    assert client.get("/cabinet/superadmin/payments/export?month=2026-10", headers=json_headers).status_code == 403
    assert client.get("/cabinet/superadmin/payment-import", headers=json_headers).status_code == 403
    assert client.post("/cabinet/superadmin/payments/mark", headers=json_headers, data={
        "user_id": student.id, "month": "2026-10", "amount": "1", "paid_on": "2026-10-12",
    }).status_code == 403
    assert client.post(f"/cabinet/superadmin/users/{student.id}/payment", headers=json_headers, data={
        "pay_window_start": "10", "pay_window_end": "15", "pay_cohort": "", "pay_price": "", "paid_month": "",
    }).status_code == 403
    profile = client.get(f"/cabinet/students/{student.id}/profile")
    if profile.status_code == 200:
        assert profile.json()["student"]["manage"].get("payment") is None
    assert db.query(Payment).count() == 0


def test_named_moderator_has_everything_about_payments(db, client, session_factory, prices, make_payer, user_factory, monkeypatch):
    """Владелец 14.10.2026: «открой доступ только Александрии» — модератору
    поимённо (`PAYMENTS_EXTRA_USER_IDS`) всё, как у ГП; другим модераторам — нет."""
    named = user_factory(vk_id=990_070, name="Александрия", role_name="модератор")
    other = user_factory(vk_id=990_071, name="Маркетолог", role_name="модератор")
    monkeypatch.setattr(ps, "PAYMENTS_EXTRA_USER_IDS", frozenset({named.id}))
    student = make_payer("Иванова Анна")
    json_headers = {"Accept": "application/json"}

    client.cookies.set("session_id", session_factory(named).id)
    assert client.get("/cabinet/superadmin/payments?month=2026-10").status_code == 200
    assert client.get("/cabinet/superadmin/payments/export?month=2026-10").status_code == 200
    assert client.get("/cabinet/superadmin/payment-import").status_code == 200
    manage = client.get(f"/cabinet/students/{student.id}/profile").json()["student"]["manage"]
    assert manage["payment"] is not None and manage["payment_can_edit"] is True
    assert 'href="/cabinet/superadmin/payments"' in client.get("/cabinet/students").text
    saved = client.post(f"/cabinet/superadmin/users/{student.id}/payment", headers=json_headers, data={
        "pay_window_start": "25", "pay_window_end": "30", "pay_cohort": "", "pay_price": "13255", "paid_month": "",
    })
    assert saved.status_code == 200, saved.text
    marked = client.post("/cabinet/superadmin/payments/mark", headers=json_headers, data={
        "user_id": student.id, "month": "2026-10", "amount": "1", "paid_on": "2026-10-12",
    })
    assert marked.status_code == 200, marked.text

    client.cookies.set("session_id", session_factory(other).id)
    assert 'href="/cabinet/superadmin/payments"' not in client.get("/cabinet/students").text
    assert client.get("/cabinet/superadmin/payments?month=2026-10", headers=json_headers).status_code == 403
    assert client.get("/cabinet/superadmin/payment-import", headers=json_headers).status_code == 403


def test_superadmin_sees_payment_tabs(db, client, session_factory, people):
    client.cookies.set("session_id", session_factory(people["chief"]).id)

    page = client.get("/cabinet/superadmin/users")

    assert "/cabinet/superadmin/payments" in page.text and "/cabinet/superadmin/payment-import" in page.text
