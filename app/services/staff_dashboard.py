"""Shared aggregates for the Chief Teacher and Superadmin dashboards."""

import io
import re
from datetime import date, datetime, timedelta, timezone
from statistics import median
from typing import TypedDict

from sqlalchemy import func, or_
from sqlalchemy.orm import Session as DBSession, aliased

from app.models.activity_event import StudentActivityEvent
from app.models.learning_topic import LearningTopic
from app.models.learning_video import LearningVideo
from app.models.video_progress import VideoProgress
from app.models.task_block import BLOCK_TIMED, BLOCK_VIDEO, TaskBlock, TaskBlockState, TaskBlockSubmission, SUBMISSION_BLOCK_TYPES
from app.models.tracker import ITEM_MOCK_EXAM, STATUS_DONE, TrackerTask
from app.constants import (
    MOCK_SUBJECTS, PROGRAM_HOLIDAY, REPORT_EXCLUDED_USER_IDS, STUDY_MODE_LABELS, TARIFF_CONFIDENT_MAX,
    TARIFF_DISPLAY, TARIFF_SELF, TARIFF_WITH_YOU, TARIFFS_CURRENT,
)
from app.models.role import Role
from app.models.session import Session as UserSession
from app.models.user import User
from app.models.work import Work, WORK_TYPE_AFTER, WORK_TYPE_BEFORE
from app.services.point_a import student_point_a_level
from app.services.report_scope import current_report_program, reportable_student_ids, reportable_students_q
from app.services.submission_edit import upload_deadline
from app.services.task_blocks import (
    BlockViewer, completed_after_deadline, feed_visible_blocks, get_audiences,
    get_submit_deadlines, get_task_submit_deadlines,
)
from app.services.tracker import (
    cycle_deadline_lookup, cycle_label, program_learners, program_students, stage_cycle_ordinal,
    stage_sibling_cycles, task_audience_user_ids,
)
from app.services.tz import msk_date_text, msk_midnight, msk_text


REGISTRATION_STATS_SINCE = date(2026, 9, 19)
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def parse_registration_date(raw: str | None, fallback: date | None = None) -> date | None:
    try:
        return date.fromisoformat(raw) if raw else fallback
    except ValueError:
        return fallback


class TariffRegistrationItem(TypedDict):
    tariff: str
    label: str
    count: int


class TariffRegistrationStats(TypedDict):
    since_label: str
    by_tariff: list[TariffRegistrationItem]
    without_tariff: int
    total: int
    students: list["TariffRegistrationStudent"]
    period_from: str
    period_to: str
    tariff_filter: str


class TariffRegistrationStudent(TypedDict):
    id: int
    name: str
    username: str
    tariff: str
    tariff_label: str
    created_at: datetime


def get_tariff_registration_stats(
    db: DBSession,
    *,
    period_from: date | None = None,
    period_to: date | None = None,
    tariff_filter: str = "",
) -> TariffRegistrationStats:
    """Count student registrations since tariff tracking started in Moscow time.

    Считаются только ученики в учёте (`report_scope`, владелец 07.10.2026):
    заблокированные, архивные, удалённые и с истёкшей подпиской выпадают и из
    прошлых регистраций. Сотрудники сюда не попадают никогда.
    """
    period_from = period_from or REGISTRATION_STATS_SINCE
    query = (
        db.query(
            User.id,
            User.name,
            User.first_name,
            User.last_name,
            User.tg_username,
            User.tariff,
            User.created_at,
        )
        .join(Role, User.role_id == Role.id)
        .filter(
            Role.rank == 1,
            User.created_at >= msk_midnight(period_from),
            User.id.in_(reportable_student_ids(db)),
        )
    )
    if period_to:
        query = query.filter(User.created_at < msk_midnight(period_to + timedelta(days=1)))
    if tariff_filter in TARIFFS_CURRENT:
        query = query.filter(User.tariff == tariff_filter)
    elif tariff_filter == "__none__":
        query = query.filter(or_(User.tariff.is_(None), User.tariff == ""))
    rows = query.order_by(User.created_at.desc(), User.id.desc()).all()
    raw_counts = {tariff: 0 for tariff in TARIFFS_CURRENT}
    without_tariff = 0
    students: list[TariffRegistrationStudent] = []
    for row in rows:
        tariff = (row.tariff or "").strip()
        if tariff in raw_counts:
            raw_counts[tariff] += 1
            tariff_label = TARIFF_DISPLAY[tariff]
        elif not tariff:
            without_tariff += 1
            tariff_label = "Без тарифа"
        else:
            # Неизвестный или старый тариф не превращаем в один из трёх
            # действующих. Запись остаётся в списке для ручной проверки.
            tariff_label = tariff

        username = (row.tg_username or "").strip().lstrip("@")
        students.append(
            {
                "id": row.id,
                "name": f"{row.last_name or ''} {row.first_name or row.name}".strip(),
                "username": f"@{username}" if username else "Не указан",
                "tariff": tariff,
                "tariff_label": tariff_label,
                "created_at": row.created_at,
            }
        )

    students.sort(key=lambda student: (_tariff_rank(student["tariff"]), student["tariff_label"].casefold(), student["name"].casefold()))

    by_tariff = [
        {
            "tariff": tariff,
            "label": TARIFF_DISPLAY[tariff],
            "count": raw_counts.get(tariff, 0),
        }
        for tariff in TARIFFS_CURRENT
    ]
    return {
        "since_label": period_from.strftime("%d.%m.%Y"),
        "by_tariff": by_tariff,
        "without_tariff": without_tariff,
        "total": sum(item["count"] for item in by_tariff) + without_tariff,
        "students": students,
        "period_from": period_from.isoformat(),
        "period_to": period_to.isoformat() if period_to else "",
        "tariff_filter": tariff_filter,
    }


# Порядок тарифов в списке и выгрузке — служба заботы 10.10.2026: «сначала все
# уверенный максимум, потом я с вами, потом я сам». Отработавшие — следом,
# без тарифа — в конце.
_TARIFF_ORDER = (TARIFF_CONFIDENT_MAX, TARIFF_WITH_YOU, TARIFF_SELF)


def _tariff_rank(tariff: str) -> int:
    if tariff in _TARIFF_ORDER:
        return _TARIFF_ORDER.index(tariff)
    return len(_TARIFF_ORDER) + (1 if not tariff else 0)


def _timed_scores_by_subject(db: DBSession, student_ids: list[int]) -> dict[int, dict[str, int]]:
    """Средний балл за контрольные на время по предметам, на ученика.

    Та же выборка, что у шапки карточки (`stats.avg_score_by_subject_all_time`),
    сужена до блоков `timed`: служба заботы 11.10.2026 просила «оценку за
    контрольные». Предмет — у блока, если пусто — у задания.
    """
    if not student_ids:
        return {}
    rows = (
        db.query(TaskBlockSubmission.user_id, TaskBlock.subject, TrackerTask.subject, TaskBlockSubmission.score)
        .join(TaskBlock, TaskBlockSubmission.block_id == TaskBlock.id)
        .join(TrackerTask, TaskBlock.task_id == TrackerTask.id)
        .filter(
            TaskBlockSubmission.user_id.in_(student_ids),
            TaskBlockSubmission.score.isnot(None),
            TaskBlock.block_type == BLOCK_TIMED,
            TrackerTask.deleted_at.is_(None),
        )
        .all()
    )
    collected: dict[int, dict[str, list[float]]] = {}
    for user_id, block_subject, task_subject, score in rows:
        subject = block_subject or task_subject
        if subject in MOCK_SUBJECTS:
            collected.setdefault(user_id, {}).setdefault(subject, []).append(float(score))
    return {
        user_id: {subject: round(sum(values) / len(values)) for subject, values in by_subject.items()}
        for user_id, by_subject in collected.items()
    }


# (заголовок, ширина колонки)
_REGISTRATION_XLSX_COLUMNS = (
    ("Фамилия Имя", 30),
    ("Ник Telegram", 22),
    ("Тариф", 20),
    ("Год поступления в вуз", 12),
    ("Р — средний балл", 10),
    ("К — средний балл", 10),
    ("Р — контрольные", 10),
    ("К — контрольные", 10),
    ("Точка А — балл", 10),
    ("Точка А — уровень", 10),
    ("Точка А — оценено", 12),
    ("Стоимость в месяц, ₽", 14),
    ("Набор", 14),
    ("Оплачено до", 12),
    ("Начало обучения", 12),
    ("Регистрация", 20),
    ("Последний вход", 20),
    ("Куратор", 24),
    ("Формат", 10),
    ("Предметы экзамена", 20),
    ("ID", 7),
)


def build_tariff_registration_xlsx(db: DBSession, stats: TariffRegistrationStats) -> bytes:
    """Excel-выгрузка списка «Регистрации по тарифам»: всё про учёбу ученика.

    Служба заботы 10–11.10.2026: «имя фамилия, ник, тариф, год поступления,
    балл за Р, за К, точка А — и я внесу стоимость, какую платят ежемесячно».
    Строки и их порядок — те же, что в списке на странице. Баллы заново не
    считаются: Р/К — как в шапке карточки (`stats.avg_score_by_subject_all_time`),
    точка А — `point_a.student_point_a`, балл и уровень только у разобранной
    целиком, как в карточке. Цена — `payments.price_kop` (своя или из
    справочника), пусто — вписать руками.

    Контактов нет намеренно — ни телефонов, ни почты, ни данных родителя,
    ни города и даты рождения (владелец 11.10.2026: «оставь всё, что
    касается учебного процесса»): файл уходит из кабинета по чатам.
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from app.models.payment import COHORT_LABELS
    from app.services.payments import price_kop
    from app.services.point_a import point_a_level, student_point_a
    from app.services.stats import avg_score_by_subject_all_time

    student_ids = [student["id"] for student in stats["students"]]
    users = {user.id: user for user in db.query(User).filter(User.id.in_(student_ids))} if student_ids else {}
    curator_ids = {user.curator_id for user in users.values() if user.curator_id}
    curators = {
        curator.id: f"{curator.last_name or ''} {curator.first_name or curator.name or ''}".strip()
        for curator in db.query(User).filter(User.id.in_(curator_ids))
    } if curator_ids else {}
    timed = _timed_scores_by_subject(db, student_ids)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ученики"
    ws.append([title for title, _ in _REGISTRATION_XLSX_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, (_, width) in enumerate(_REGISTRATION_XLSX_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(index)].width = width

    for student in stats["students"]:
        user = users.get(student["id"])
        if user is None:
            continue
        scores = avg_score_by_subject_all_time(db, user.id)
        controls = timed.get(user.id, {})
        point_a = student_point_a(db, user, with_images=False)
        point_a_score = point_a.average if point_a.is_done else None
        price = price_kop(db, user)
        ws.append([
            student["name"],
            student["username"],
            student["tariff_label"],
            user.university_year,
            scores.get("Рисунок"),
            scores.get("Композиция"),
            controls.get("Рисунок"),
            controls.get("Композиция"),
            point_a_score,
            point_a_level(point_a_score) if point_a_score is not None else None,
            f"{point_a.scored_count} из {len(point_a.plates)}" if point_a.plates else "нет работ",
            price / 100 if price else None,
            COHORT_LABELS.get(user.pay_cohort or "", ""),
            msk_date_text(user.paid_until),
            user.enrolled_at.strftime("%m.%Y") if user.enrolled_at else "",
            msk_text(user.created_at),
            msk_text(user.last_login_at),
            curators.get(user.curator_id, ""),
            STUDY_MODE_LABELS.get(user.study_mode or "", ""),
            user.exam_subjects or "",
            user.id,
        ])

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


_ACTIVITY_LABELS = {
    "login": "Вход в кабинет",
    "portfolio_upload": "Загрузка портфолио",
    "work_upload": "Загрузка работы",
}


def _assignment_activity(db: DBSession, students: list[User], *, theory: bool = False) -> list[dict]:
    """Открытые блоки сдачи работы (домашка и контрольная на время): кому
    положено сдавать, кто сдал, а у несдавших — прошёл ли срок.

    `theory=True` — то же по видео-блокам (Александрия 22.10.2026: «Теория —
    месяц, цикл, видео, и список, кто посмотрел и дата просмотра»).
    «Посмотрел» — видео-блок закрыт: кружок ученика или зачёт преподавателя
    (`video_credit`), тот же `TaskBlockState`, что «сдал» у работ. Срока у
    видео нет — все непосмотревшие в `pending`. Пробник не считается, как в
    контроле просмотра. Непосмотревшему подпись говорит, сколько он уже
    досмотрел: «Смотрел 27 из 40 мин» (Александрия 23.10.2026) — покрытие
    текущего прохода `VideoProgress.covered_seconds`, то же, по которому
    засчитывается просмотр; такие ученики ещё и в `started`.

    У каждого задания ещё `topic_id` (цикл — для выбора «месяц → цикл») и
    `done_at` — когда сдал или посмотрел, по Москве.

    Фильтр «Задание → Сдали / Не сдали» в «Учениках поимённо» и копирование
    списка (владелец 03.10.2026: «выбрал задание, фильтр сдал или не сдал, и
    показался список, который можно скопировать»).

    Правила не свои — те же, что у напоминаний о сроке
    (`student_reminders._collect_deadlines`): кому видно задание —
    `tracker.task_audience_user_ids`, с кого спрашивают —
    `program_students`/`program_learners`, блок по тарифу —
    `feed_visible_blocks`, сдал — `TaskBlockState` в статусе «сделано», срок —
    `submission_edit.upload_deadline`. До 03.10.2026 здесь жила своя копия
    адресации, и в «не сдали» по каждой контрольной попадали шестеро
    пробников, у которых доступ кончился 27.09: бот им уже не напоминал,
    а статистика числила должниками.
    """
    now = datetime.now(timezone.utc)
    kind_filter = (
        (TaskBlock.block_type == BLOCK_VIDEO, TaskBlock.video_id.isnot(None),
         TrackerTask.kind != ITEM_MOCK_EXAM)
        if theory else (TaskBlock.block_type.in_(SUBMISSION_BLOCK_TYPES),)
    )
    candidates = (
        db.query(TaskBlock, TrackerTask)
        .join(TrackerTask, TaskBlock.task_id == TrackerTask.id)
        .outerjoin(LearningTopic, TrackerTask.topic_id == LearningTopic.id)
        .filter(
            *kind_filter,
            TrackerTask.is_published.is_(True),
            TrackerTask.deleted_at.is_(None),
            or_(TrackerTask.starts_at.is_(None), TrackerTask.starts_at <= now),
            or_(TaskBlock.opens_at.is_(None), TaskBlock.opens_at <= now),
            or_(TrackerTask.topic_id.is_(None),
                (LearningTopic.is_published.is_(True)) &
                (LearningTopic.deleted_at.is_(None)) &
                (LearningTopic.opens_at <= now)),
        )
        .order_by(TrackerTask.id.desc(), TaskBlock.sort_order, TaskBlock.id)
        .all()
    )
    if not candidates:
        return []

    learners = program_learners(program_students(db, now))
    # Строки таблицы — свои у сводки; адресат без строки в списке не нужен.
    reachable = learners.keys() & {student.id for student in students}
    block_ids = [block.id for block, _ in candidates]
    task_ids = list({task.id for _, task in candidates})
    audiences = get_audiences(db, block_ids)
    # Уровень точки А считается лениво и один раз на ученика (`BlockViewer`).
    viewers = {
        uid: BlockViewer(db, user_id=uid, tariff=learners[uid].tariff) for uid in reachable
    }
    block_deadlines = get_submit_deadlines(db, block_ids)
    task_deadlines = get_task_submit_deadlines(db, task_ids)
    # Срок цикла — запасной срок отметки «Сдал после срока» (06.10.2026).
    cycle_deadline = cycle_deadline_lookup(db, {task.topic_id for _, task in candidates})
    audience = {task_id: task_audience_user_ids(db, task_id) & reachable for task_id in task_ids}
    done = {
        (state.block_id, state.user_id): state
        for state in db.query(TaskBlockState).filter(
            TaskBlockState.block_id.in_(block_ids),
            TaskBlockState.status == STATUS_DONE,
        )
    }

    videos = {
        video.id: video
        for video in db.query(LearningVideo).filter(
            LearningVideo.id.in_({block.video_id for block, _ in candidates})
        )
    } if theory else {}
    video_titles = {video_id: video.title for video_id, video in videos.items()}
    # Сколько досмотрено у непосмотревших — одним запросом на всю выборку.
    progress: dict[tuple[int, str], VideoProgress] = {}
    if theory and reachable:
        for row in db.query(VideoProgress).filter(
            VideoProgress.video_id.in_({video.bunny_video_id for video in videos.values()}),
            VideoProgress.user_id.in_(reachable),
        ):
            progress[(row.user_id, row.video_id)] = row

    assignments = []
    for block, task in candidates:
        eligible = sorted(
            uid for uid in audience[task.id]
            if feed_visible_blocks([block], audiences, viewers[uid])
        )
        submitted, overdue, pending, started, notes, done_at = [], [], [], [], {}, {}
        for uid in eligible:
            tariff = learners[uid].tariff
            state = done.get((block.id, uid))
            if state is not None:
                submitted.append(uid)
                done_at[uid] = msk_text(state.completed_at or state.updated_at)
            if theory:
                if state is not None:
                    notes[uid] = "Посмотрел"
                    continue
                pending.append(uid)
                video = videos.get(block.video_id)
                watched = _watched_note(
                    progress.get((uid, video.bunny_video_id)) if video else None, video,
                )
                if watched:
                    started.append(uid)
                notes[uid] = watched or "Не посмотрел"
                continue
            if state is not None:
                late = completed_after_deadline(
                    block, task, state, user_tariff=tariff,
                    block_overrides=block_deadlines.get(block.id),
                    task_overrides=task_deadlines.get(task.id),
                    cycle_deadline=cycle_deadline(task.topic_id, tariff),
                )
                notes[uid] = "Сдал после срока" if late else "Сдал"
                continue
            deadline = upload_deadline(
                task, block, user_tariff=tariff,
                tariff_deadlines=block_deadlines.get(block.id),
                task_tariff_deadlines=task_deadlines.get(task.id),
            )
            if deadline is not None and deadline <= now:
                overdue.append(uid)
                notes[uid] = f"Не сдал · срок прошёл {msk_text(deadline)}"
            else:
                pending.append(uid)
                notes[uid] = (
                    f"Не сдал · срок до {msk_text(deadline)}" if deadline is not None
                    else "Не сдал · срок не задан"
                )
        if theory:
            # Видео узнают по названию ролика; задание — уточнение, если
            # называется иначе (в одном задании бывает несколько роликов).
            label = (video_titles.get(block.video_id) or "").strip() or task.title
            if label.casefold() != task.title.casefold():
                label += f" · {task.title}"
        else:
            label = task.title
            if block.title and block.title.strip() and block.title.strip().casefold() != task.title.casefold():
                label += f" · {block.title.strip()}"
            # Предмет — подпись в выпадающем списке, на подсчёт не влияет. Стоит
            # первым: на телефоне длинное название обрезается справа.
            subject = block.subject or task.subject
            if subject:
                label = f"{subject}: {label}"
        assignments.append({
            "id": block.id,
            "label": label,
            "topic_id": task.topic_id,
            # Порядок внутри цикла — как в ленте ученика.
            "order": [task.sort_order or 0, task.id, block.sort_order or 0, block.id],
            "eligible": eligible,
            "submitted": submitted,
            "overdue": overdue,
            "pending": pending,
            "started": started,
            "notes": notes,
            "done_at": done_at,
        })
    return assignments


def _watched_note(progress: VideoProgress | None, video: LearningVideo | None) -> str | None:
    """«Смотрел 27 из 40 мин» — непосмотревшему, который начал; None — не
    начинал. Длина — у ролика, а если её ещё нет, из плеера (`VideoProgress`)."""
    if progress is None or not progress.covered_seconds:
        return None
    duration = (video.duration_seconds if video else None) or progress.duration_seconds
    minutes = int(progress.covered_seconds // 60)
    seen = str(minutes) if minutes else "меньше минуты"
    if duration:
        return f"Смотрел {seen} из {max(1, round(duration / 60))} мин"
    return f"Смотрел {seen} мин" if minutes else "Смотрел меньше минуты"


def _program_months(db: DBSession, items: list[dict]) -> list[dict]:
    """Выбор «месяц → цикл» для «Теории» и «Практики» в «Учениках поимённо»
    (Александрия 22.10.2026: «выбор месяца, выбор цикла 1/2/3/4 неделя»).

    Месяц — этап программы (`LearningTopic kind='stage'`), цикл — его цикл с
    подписью `tracker.cycle_label` («Цикл N», как у ученика). Берутся только
    циклы, в которых есть задания из `items`: пустой выбор никому не нужен.
    Цикл без этапа — в «Без месяца», задание без цикла — в «Без цикла»
    (`topic_id` пуст, ключ цикла 0). Порядок — по дате открытия, свежие
    месяцы сверху: смотрят обычно текущий.
    """
    topic_ids = {item["topic_id"] for item in items if item["topic_id"]}
    cycles = (
        db.query(LearningTopic).filter(LearningTopic.id.in_(topic_ids)).all()
        if topic_ids else []
    )
    stage_ids = {cycle.parent_id for cycle in cycles if cycle.parent_id}
    stages = {
        stage.id: stage
        for stage in db.query(LearningTopic).filter(LearningTopic.id.in_(stage_ids))
    } if stage_ids else {}
    months: dict[int, dict] = {}
    for cycle in sorted(cycles, key=lambda c: (c.opens_at, c.sort_order, c.id)):
        stage = stages.get(cycle.parent_id)
        key = stage.id if stage else 0
        month = months.setdefault(key, {
            "id": key,
            "label": stage.title if stage else "Без месяца",
            "opens_at": (stage.opens_at if stage else cycle.opens_at),
            "cycles": [],
        })
        month["cycles"].append({"id": cycle.id, "label": _numbered_cycle_label(db, cycle)})
    if any(not item["topic_id"] for item in items):
        months.setdefault(-1, {"id": -1, "label": "Без цикла", "opens_at": None, "cycles": []})
        months[-1]["cycles"].append({"id": 0, "label": "Без цикла"})
    ordered = sorted(
        months.values(),
        key=lambda m: (m["opens_at"] is not None, m["opens_at"] or datetime.min), reverse=True,
    )
    # Без цикла и без месяца — в конец, после настоящих месяцев.
    ordered.sort(key=lambda m: m["id"] <= 0)
    for month in ordered:
        month.pop("opens_at")
    return ordered


def _numbered_cycle_label(db: DBSession, cycle: LearningTopic) -> str:
    """Подпись цикла в выборе статистики с номером по порядку в этапе:
    «5. Контрольная работа» (Александрия 23.10.2026: «писать номер рядом с
    названием цикла, чтобы это отражало последовательность»). Номер уже в
    названии («Цикл 1», «1 неделя») — подпись как есть, без «1. Цикл 1».
    Единственный цикл этапа номера не получает: «Портфолио», а не
    «1. Портфолио» — последовательности там нет.
    Номер тот же, что даёт безымянному циклу «Цикл N» (`stage_cycle_ordinal`)."""
    label = cycle_label(db, cycle)
    ordinal = stage_cycle_ordinal(db, cycle) if cycle.parent_id else None
    if ordinal is None or re.search(rf"(?<!\d){ordinal}(?!\d)", label):
        return label
    if ordinal == 1 and len(stage_sibling_cycles(db, cycle.parent_id)) == 1:
        return label
    return f"{ordinal}. {label}"


def get_student_activity_overview(
    db: DBSession, event_limit: int = 200, *, include_assignments: bool = False,
    curator_id: int | None = None,
) -> dict:
    """Return per-student lifecycle metrics and the append-only action journal.

    `curator_id` — «Ученики поимённо» куратора: только его ученики
    (`User.curator_id`, то же правило, что у его списка «Учеников»),
    Александрия 22.10.2026: «чтобы видели ситуацию только по своей группе».

    С `include_assignments` ещё «Теория» (`videos` — видео-блоки), выбор
    «месяц → цикл» (`months`) и уровень точки А у каждого ученика.
    """
    query = reportable_students_q(db)
    if curator_id is not None:
        query = query.filter(User.curator_id == curator_id)
    students = query.order_by(User.last_name, User.first_name, User.id).all()
    student_ids = [student.id for student in students]
    if not student_ids:
        return {"students": [], "events": [], "assignments": [], "scores": [],
                "videos": [], "months": [], "tariffs": []}

    login_rows = (
        db.query(
            UserSession.user_id,
            func.min(UserSession.created_at).label("first_login"),
            func.max(UserSession.created_at).label("last_login"),
            func.count(UserSession.id).label("login_count"),
        )
        .filter(UserSession.user_id.in_(student_ids), UserSession.impersonated_by_id.is_(None))
        .group_by(UserSession.user_id)
        .all()
    )
    upload_rows = (
        db.query(
            Work.user_id,
            func.count(Work.id).label("upload_count"),
            func.min(Work.created_at).label("first_upload"),
            func.max(Work.created_at).label("last_upload"),
        )
        .filter(
            Work.user_id.in_(student_ids),
            Work.status == "success",
            Work.work_type.in_((WORK_TYPE_BEFORE, WORK_TYPE_AFTER)),
        )
        .group_by(Work.user_id)
        .all()
    )
    login_by_user = {row.user_id: row for row in login_rows}
    upload_by_user = {row.user_id: row for row in upload_rows}
    portfolio_before_user_ids = {
        row.user_id
        for row in (
            db.query(Work.user_id)
            .filter(
                Work.user_id.in_(student_ids),
                Work.status == "success",
                Work.work_type == WORK_TYPE_BEFORE,
            )
            .distinct()
            .all()
        )
    }
    events = (
        db.query(StudentActivityEvent, User)
        .join(User, StudentActivityEvent.user_id == User.id)
        .filter(StudentActivityEvent.user_id.in_(student_ids))
        .order_by(StudentActivityEvent.created_at.desc(), StudentActivityEvent.id.desc())
        .limit(event_limit)
        .all()
    )
    latest_event_by_user: dict[int, StudentActivityEvent] = {}
    for event, _user in events:
        latest_event_by_user.setdefault(event.user_id, event)

    overview = []
    for student in students:
        login = login_by_user.get(student.id)
        upload = upload_by_user.get(student.id)
        latest_event = latest_event_by_user.get(student.id)
        overview.append({
            "id": student.id,
            "name": f"{student.last_name or ''} {student.first_name or student.name}".strip(),
            "username": f"@{(student.tg_username or '').strip().lstrip('@')}" if student.tg_username else "Не указан",
            "tariff": student.tariff or "Без тарифа",
            "registered_at": student.created_at,
            "first_login": login.first_login if login else None,
            "last_login": login.last_login if login else None,
            "login_count": int(login.login_count) if login else 0,
            "upload_count": int(upload.upload_count) if upload else 0,
            "has_portfolio_before": student.id in portfolio_before_user_ids,
            "first_upload": upload.first_upload if upload else None,
            "last_upload": upload.last_upload if upload else None,
            "last_action_at": latest_event.created_at if latest_event else (login.last_login if login else None),
            "last_action": _ACTIVITY_LABELS.get(latest_event.event_type, latest_event.event_type) if latest_event else ("Вход в кабинет" if login else "Нет действий"),
        })

    journal = [
        {
            "created_at": event.created_at,
            "student_name": f"{user.last_name or ''} {user.first_name or user.name}".strip(),
            "username": f"@{(user.tg_username or '').strip().lstrip('@')}" if user.tg_username else "Не указан",
            "event_type": event.event_type,
            "action": _ACTIVITY_LABELS.get(event.event_type, event.event_type),
            "details": event.details,
        }
        for event, user in events
    ]
    if not include_assignments:
        return {"students": overview, "events": journal, "assignments": [], "scores": [],
                "videos": [], "months": [], "tariffs": []}
    assignments = _assignment_activity(db, students)
    videos = _assignment_activity(db, students, theory=True)
    # В выбор «месяц → цикл → задание» — только то, что положено хоть одному
    # ученику выборки: в каникулярном режиме иначе шли все задания годового
    # курса (23.10.2026). Баллы считают по полному списку: сдавший сохраняет
    # балл, даже если задание ему больше не адресовано.
    shown = [item for item in assignments if item["eligible"]]
    shown_videos = [item for item in videos if item["eligible"]]
    holiday = current_report_program() == PROGRAM_HOLIDAY
    if holiday:
        # У каникулярного тарифа нет — фильтр идёт по треку, тегу группы
        # (Александрия 23.10.2026: «5 линий… отслеживать с 25 по 1 ноября»).
        from app.services.intake_link import tracks_by_user

        tracks = tracks_by_user(db, student_ids)
        for row in overview:
            titles = [link.title for link in tracks.get(row["id"], [])]
            row["tariff"] = " + ".join(titles) if titles else "Без трека"
        groups = sorted({row["tariff"] for row in overview}, key=lambda t: (t == "Без трека", t))
    else:
        # Уровень точки А — фильтр «Уровень» (1 / 2 / ещё не определён).
        levels = {student.id: student_point_a_level(db, student) for student in students}
        for row in overview:
            row["level"] = levels.get(row["id"])
        # Фильтр «Тариф» — в порядке службы заботы (`_TARIFF_ORDER`).
        groups = sorted({row["tariff"] for row in overview},
                        key=lambda t: (_tariff_rank("" if t == "Без тарифа" else t), t))
    return {"students": overview, "events": journal, "assignments": shown,
            "scores": _assignment_scores(db, student_ids, assignments),
            "videos": shown_videos, "months": _program_months(db, shown + shown_videos),
            "tariffs": groups, "group_label": "Трек" if holiday else "Тариф",
            "show_level": not holiday}


def _score_value(score) -> int | float:
    """`55.00` → 55, `57.50` → 57.5: в списке, сводке и копии без хвоста нулей."""
    value = float(score)
    return int(value) if value.is_integer() else value


def _assignment_scores(db: DBSession, student_ids: list[int], assignments: list[dict]) -> list[dict]:
    """Баллы за сдачи по заданию — карточка «Баллы за задания» на «Статистике
    активности» (владелец 03.10.2026: «сколько людей на 55, сколько на 65» по
    контрольной).

    Задания и подписи — из `_assignment_activity`, своей выборки «какие
    задания открыты» здесь нет. Считаются все сдавшие из строк сводки, без
    фильтра `eligible`. Строки сводки — ученики в учёте (`report_scope`), так
    что с 07.10.2026 балл ученика с истёкшей подпиской сюда не входит
    (владелец: «только с активной подпиской»). Само распределение по
    баллам строит страница из `scored` — так же, как разбивку по тарифам
    в «Учениках поимённо». Задание без единого балла не отдаётся.
    """
    if not assignments or not student_ids:
        return []
    scorer_user = aliased(User)
    rows = (
        db.query(TaskBlockSubmission, scorer_user)
        .outerjoin(scorer_user, TaskBlockSubmission.scored_by_id == scorer_user.id)
        .filter(
            TaskBlockSubmission.block_id.in_([item["id"] for item in assignments]),
            TaskBlockSubmission.user_id.in_(student_ids),
            TaskBlockSubmission.submitted_at.isnot(None),
        )
        .all()
    )
    by_block: dict[int, list[tuple[TaskBlockSubmission, User | None]]] = {}
    for submission, scorer in rows:
        by_block.setdefault(submission.block_id, []).append((submission, scorer))

    result = []
    for item in assignments:
        scored, scorers, unscored = {}, {}, []
        for submission, scorer in by_block.get(item["id"], []):
            if submission.score is None:
                unscored.append(submission.user_id)
                continue
            scored[submission.user_id] = _score_value(submission.score)
            if scorer is not None:
                scorers[submission.user_id] = (
                    f"{scorer.last_name or ''} {scorer.first_name or scorer.name or ''}".strip()
                )
        if not scored:
            continue
        values = list(scored.values())
        result.append({
            "id": item["id"],
            "label": item["label"],
            "submitted": len(scored) + len(unscored),
            "scored": scored,
            "unscored": sorted(unscored),
            "scorers": scorers,
            "avg": round(sum(values) / len(values), 1),
            "median": _score_value(median(values)),
        })
    return result


def student_assignments(db: DBSession, student: User) -> list[dict]:
    """Задания со сдачей работы у одного ученика — раздел вкладки
    «Статистика» его карточки (владелец 06.10.2026: «собрать туда всё, что
    касается ребёнка из вкладки статистика»).

    Тот же расчёт, что «Ученики поимённо» и «Баллы за задания» на
    «Статистике активности»: `_assignment_activity` и `_assignment_scores`
    для списка из одного ученика. Задание, которое ему сдавать не положено,
    не отдаётся. Архивного и заблокированного `program_learners` не числит
    учениками программы — у них список пуст.
    """
    mine = [item for item in _assignment_activity(db, [student]) if student.id in item["eligible"]]
    scores = {item["id"]: item for item in _assignment_scores(db, [student.id], mine)}
    rows = []
    for item in mine:
        if student.id in item["submitted"]:
            status = "late" if item["notes"][student.id] == "Сдал после срока" else "done"
        else:
            status = "overdue" if student.id in item["overdue"] else "pending"
        score = scores.get(item["id"])
        rows.append({
            "label": item["label"],
            "status": status,
            "note": item["notes"][student.id],
            "score": score["scored"].get(student.id) if score else None,
            "scorer": score["scorers"].get(student.id) if score else None,
        })
    return rows


# ═══════════════════════════════════════════════════════════════════════════════
# Главный экран Главного преподавателя и суперадмина
# ═══════════════════════════════════════════════════════════════════════════════

_MONTHS_PREP = ["", "январе", "феврале", "марте", "апреле", "мае", "июне",
                "июле", "августе", "сентябре", "октябре", "ноябре", "декабре"]


def load_staff_dashboard(db: DBSession, user: dict, now: datetime) -> dict:
    """Данные для `cabinet_staff.html` — один экран у ГП и суперадмина.

    До 28.09.2026 жили двумя почти одинаковыми копиями в `cabinet_admin.py` и
    `cabinet_superadmin.py` и расходились по мелочам. Здесь же сняты расчёты,
    которые шаблон не показывал: регистрации и активность учеников (их
    карточки живут на «Статистике активности»), лента последних загрузок
    (ветка шаблона для ранга ниже 4, до которой ни один маршрут не доходит)
    и окно загрузки портфолио — с 09.09.2026 оно ни на что не влияет
    (`docs/invariants/portfolio.md`), а кнопка на дашборде им управляла.

    Всё учебное считается по ученикам (ранг 1): в «Учениках» и «Новых
    учениках» сотрудники не нужны. Ученики — только те, кто в учёте
    (`report_scope`: без заблокированных, удалённых, с истёкшей подпиской и
    служебных). Заблокированные показаны отдельной подписью «ещё N», в число
    учеников не входят.

    Очереди пробников на дашборде нет (владелец 29.09.2026 снял блок «Ждёт
    проверки» вместе со ссылками на проверку, билеты и статистику пробников).
    Счётчика «ученики без точки А» тоже нет: `point_a.student_point_a` ходит в
    базу по каждому ученику, на 20 учениках дашборд делал 173 запроса вместо 14
    (`test_superadmin_dashboard_does_not_scale_with_data`).
    """
    from app.models.work import WORK_TYPE_MOCK_EXAM

    month_start = datetime(now.year, now.month, 1, tzinfo=timezone.utc)
    reportable_ids = reportable_student_ids(db, now)
    students_active = reportable_students_q(db, now).count()

    role_rows = (
        db.query(Role.display_name, Role.rank, func.count(User.id).label("cnt"))
        .outerjoin(User, (User.role_id == Role.id) & (User.is_active == True))  # noqa: E712
        .group_by(Role.id, Role.display_name, Role.rank)
        .order_by(Role.rank)
        .all()
    )
    # Строку учеников считает общее правило, остальные роли — по `is_active`.
    role_breakdown = [
        {"name": r.display_name, "rank": r.rank, "count": students_active if r.rank == 1 else r.cnt}
        for r in role_rows
    ]

    students_blocked = (
        db.query(User.id)
        .join(Role, User.role_id == Role.id)
        .filter(
            Role.rank == 1,
            User.is_active == False,  # noqa: E712
            User.archived_at.is_(None),
            User.deleted_at.is_(None),
            User.id.notin_(REPORT_EXCLUDED_USER_IDS),
        )
        .count()
    )
    new_students_month = (
        reportable_students_q(db, now).filter(User.created_at >= month_start).count()
    )

    works = db.query(Work).filter(Work.status == "success", Work.user_id.in_(reportable_ids))
    works_by_type = dict(
        works.with_entities(Work.work_type, func.count(Work.id)).group_by(Work.work_type).all()
    )
    total_works = sum(works_by_type.values())
    works_this_month = works.filter(Work.created_at >= month_start).count()

    avg_raw = (
        works.filter(Work.work_type == WORK_TYPE_MOCK_EXAM, Work.score.isnot(None))
        .with_entities(func.avg(Work.score))
        .scalar()
    )
    avg_mock_score = round(float(avg_raw)) if avg_raw is not None else None

    StudentAlias = aliased(User)
    curator_rows = (
        db.query(
            User.id, User.first_name, User.last_name, User.name, User.photo_url,
            func.count(StudentAlias.id).label("student_count"),
        )
        .join(Role, User.role_id == Role.id)
        .outerjoin(
            StudentAlias,
            (StudentAlias.curator_id == User.id)
            & StudentAlias.id.in_(reportable_ids),
        )
        .filter(Role.rank == 2, User.is_active == True)  # noqa: E712
        .group_by(User.id, User.first_name, User.last_name, User.name, User.photo_url)
        .order_by(func.count(StudentAlias.id).desc())
        .limit(100)
        .all()
    )
    curators = [
        {
            "id": r.id,
            "name": f"{r.last_name or ''} {r.first_name or r.name}".strip(),
            "photo_url": r.photo_url,
            "student_count": r.student_count,
        }
        for r in curator_rows
    ]

    # Модератор (ранг 3) в списке рядом с ГП: уровень у него тот же
    # (rbac.py::effective_role_rank), подпись — своя роль.
    admin_rows = (
        db.query(User.id, User.first_name, User.last_name, User.name, User.photo_url, Role.display_name)
        .join(Role, User.role_id == Role.id)
        .filter(Role.rank.in_([3, 4]), User.is_active == True)  # noqa: E712
        .order_by(User.last_name, User.first_name)
        .limit(50)
        .all()
    )
    admins = [
        {
            "id": r.id,
            "name": f"{r.last_name or ''} {r.first_name or r.name}".strip(),
            "photo_url": r.photo_url,
            "role_label": r.display_name,
        }
        for r in admin_rows
    ]

    return {
        "role_breakdown": role_breakdown,
        "students_active": students_active,
        "students_blocked": students_blocked,
        "new_students_month": new_students_month,
        "works_by_type": works_by_type,
        "total_works": total_works,
        "works_this_month": works_this_month,
        "avg_mock_score": avg_mock_score,
        "curators": curators,
        "admins": admins,
        "month_name": _MONTHS_PREP[now.month],
    }
