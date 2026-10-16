"""Экран «Оплаты»: кто за месяц оплатил, кто нет (шаг 3 плана оплаты), и
статистика оплат — итоги по тарифам, долг, сводка по месяцам, Excel
(`plans/2026-10-13-payment-stats.md`).

План — `plans/2026-09-30-apparchi-monthly-payments.md`. Заказчик: «экран кто
оплатил / кто нет, как статистика просмотров видео».

Оплачен ли месяц — по «оплачено по» ученика (`payments.paid_month`), а не по
строкам платежей: срок мог поставить человек в карточке или загрузкой
списка, и платежа за такой месяц нет. Строка платежа, если есть, даёт
подробности — сумму, дату, Продамус или вручную, кто отметил.

Деньги в итогах — только по строкам платежей. Месяц, закрытый одним сроком,
идёт отдельным счётчиком «отмечено без платежа»: при запуске 13.10.2026 всем
поставили «оплачено по сентябрь» ради блокировки, и сентябрь не должен
выглядеть как 91 оплата.

Кто на экране — те же, кто в учёте: ученики вне архива, не заблокированные,
без служебных (`REPORT_EXCLUDED_USER_IDS`). Ученики без окна оплаты стоят
отдельно: они вне автоматической оплаты, и долга у них не бывает. Так же
отдельно — ушедшие с курса (`has_left`): месяц не оплачен, а срок доступа
наступил. В должниках и «ждём оплату» их нет; оплаченный месяц ушедшего
остаётся в итогах, деньги пришли.
"""
import io
from dataclasses import dataclass, field
from datetime import date, datetime

from sqlalchemy import func
from sqlalchemy.orm import Session as DBSession

from app.constants import REPORT_EXCLUDED_USER_IDS, TARIFF_DISPLAY
from app.models.payment import (
    KIND_MONTH,
    SOURCE_MANUAL,
    STATUS_PAID,
    STATUS_REFUNDED,
    STATUS_SUM_MISMATCH,
    Payment,
)
from app.models.role import Role
from app.models.user import User
from app.services import payments
from app.services.contacts import normalize_tg_username
from app.services.student_search import student_matches
from app.services.tz import MSK_TZ, _as_utc, msk_date_text

# Состояние ученика за месяц; порядок — порядок строк на экране.
STATE_OVERDUE = "overdue"    # окно закрылось, оплаты нет
STATE_OPEN = "open"          # окно идёт
STATE_UPCOMING = "upcoming"  # окно ещё не открылось
STATE_PAID = "paid"
STATE_ORDER = (STATE_OVERDUE, STATE_OPEN, STATE_UPCOMING, STATE_PAID)
STATE_LABELS = {
    STATE_OVERDUE: "Просрочено",
    STATE_OPEN: "Окно идёт",
    STATE_UPCOMING: "Окно впереди",
    STATE_PAID: "Оплачено",
}
NO_PAYMENT_TEXT = "Отмечено без платежа"
LEFT_TEXT = "Ушёл с курса"

# Фильтр таблицы (владелец 14.10.2026: «фильтр оплатили / не оплатили»).
# «Оплатили» — месяц закрыт, в том числе сроком без платежа; «не оплатили» —
# все остальные: окно впереди, идёт или закрылось без оплаты.
FILTER_PAID = "paid"
FILTER_UNPAID = "unpaid"
FILTER_LABELS = {"": "Все", FILTER_PAID: "Оплатили", FILTER_UNPAID: "Не оплатили"}


@dataclass
class BoardRow:
    student: User
    state: str
    window_text: str
    amount_kop: int | None          # оплачено или сколько ждём
    payment: Payment | None = None  # оплаченный платёж, если есть
    detail: str = ""                # как оплачено или до какого числа окно
    problems: list[str] = field(default_factory=list)
    # Тариф строки: у оплаченного — тариф платежа (оплата могла его сменить,
    # «Тариф после оплаты»), у остальных — нынешний тариф ученика.
    tariff: str = ""
    paid_on: str = ""               # дата оплаты по Москве
    method: str = ""                # «Продамус» или «Вручную, кто отметил»

    @property
    def no_payment(self) -> bool:
        """Месяц закрыт сроком «оплачено по», денег по нему в системе нет."""
        return self.state == STATE_PAID and self.payment is None


@dataclass
class TariffTotal:
    tariff: str
    paid_count: int = 0
    paid_kop: int = 0
    no_payment_count: int = 0
    debt_count: int = 0
    debt_kop: int = 0
    waiting_count: int = 0          # окно идёт или впереди


@dataclass
class Board:
    period: date
    rows: list[BoardRow]
    no_window: list[User]
    counts: dict[str, int]
    paid_total_kop: int
    paid_count: int = 0             # оплатили деньгами
    no_payment_count: int = 0
    debt_total_kop: int = 0
    by_tariff: list[TariffTotal] = field(default_factory=list)
    left: list[User] = field(default_factory=list)   # ушли с курса, месяц не оплачен


@dataclass
class MonthTotal:
    period: date
    paid_count: int
    paid_kop: int
    no_payment_count: int
    debt_count: int
    debt_kop: int
    waiting_count: int


def parse_month(raw: str | None, now: datetime) -> date:
    """«2026-10» → 1 октября 2026; пусто или мусор — текущий месяц по Москве."""
    try:
        parsed = payments.parse_paid_month(raw)
    except payments.PaymentSettingsError:
        parsed = None
    return parsed or payments.month_start(now.astimezone(MSK_TZ).date())


def accounted_students(db: DBSession) -> list[User]:
    """Ученики учёта оплаты: вне архива, активные, без служебных. Тот же
    список читает отчёт о неоплативших (`payment_debtors`)."""
    return (
        db.query(User)
        .join(Role, Role.id == User.role_id)
        .filter(
            Role.rank == 1,
            User.deleted_at.is_(None),
            User.archived_at.is_(None),
            User.is_active == True,  # noqa: E712
            User.id.notin_(REPORT_EXCLUDED_USER_IDS),
        )
        .order_by(User.last_name, User.first_name, User.id)
        .all()
    )


def has_left(student: User, now: datetime) -> bool:
    """Ушёл с курса: срок доступа (`access_until`) поставлен и наступил.
    Владелец 16.10.2026: «если срок стоит в блокировке, значит, они ушли, и их
    не нужно в список добавлять» — уходящим закрывают доступ сроком, окно
    оплаты при этом остаётся, и без этого правила они числились бы
    должниками. Одно правило на экран «Оплаты», статистику, Excel и
    список неоплативших в боте (`payment_debtors`)."""
    return student.access_until is not None and _as_utc(student.access_until) <= now


def display_name(student: User) -> str:
    return f"{student.last_name or ''} {student.first_name or student.name}".strip()


def tariff_text(code: str | None) -> str:
    return TARIFF_DISPLAY.get(code or "", code or "–")


def contacts_text(student: User) -> str:
    """Почта и Telegram одной строкой — как в уведомлении службе заботы."""
    tg = normalize_tg_username(student.tg_username or "")
    return " · ".join(part for part in (student.email or "", f"@{tg}" if tg else "") if part)


def _msk_day(value: datetime) -> str:
    return value.astimezone(MSK_TZ).strftime("%d.%m")


def _paid_method(db: DBSession, payment: Payment) -> str:
    if payment.source == SOURCE_MANUAL:
        who = db.get(User, payment.marked_by_id) if payment.marked_by_id else None
        return f"Вручную, {display_name(who)}" if who else "Вручную"
    return "Продамус"


def month_board(db: DBSession, period: date, now: datetime,
                students: list[User] | None = None) -> Board:
    """Месяц на экране. `students` — список учёта, если он уже прочитан
    (сводка по месяцам читает его один раз)."""
    if students is None:
        students = accounted_students(db)
    by_user: dict[int, list[Payment]] = {}
    for payment in (
        db.query(Payment)
        .filter(Payment.period == period, Payment.kind == KIND_MONTH,
                Payment.user_id.in_([s.id for s in students] or [0]))
        .order_by(Payment.id)
    ):
        by_user.setdefault(payment.user_id, []).append(payment)

    rows: list[BoardRow] = []
    no_window: list[User] = []
    left: list[User] = []
    paid_total = 0
    for student in students:
        if not payments.has_window(student):
            no_window.append(student)
            continue
        own = by_user.get(student.id, [])
        paid = next((p for p in own if p.status == STATUS_PAID), None)
        problems = []
        for p in own:
            if p.status == STATUS_SUM_MISMATCH:
                problems.append(f"Пришла не та сумма: {payments.rub_text(p.paid_sum_kop or 0)} вместо {payments.rub_text(p.amount_kop)}")
            elif p.status == STATUS_REFUNDED:
                problems.append(f"Возврат {payments.rub_text(p.paid_sum_kop or p.amount_kop)}")

        paid_month = payments.paid_month(student)
        window = f"{student.pay_window_start}–{student.pay_window_end}"
        if paid is not None:
            amount = paid.paid_sum_kop or paid.amount_kop
            paid_total += amount
            paid_on = msk_date_text(paid.paid_at)
            method = _paid_method(db, paid)
            rows.append(BoardRow(
                student, STATE_PAID, window, amount, paid, f"{method}, {paid_on}", problems,
                tariff=paid.tariff or student.tariff or "", paid_on=paid_on, method=method,
            ))
            continue
        if paid_month is not None and paid_month >= period:
            rows.append(BoardRow(student, STATE_PAID, window, None, None, NO_PAYMENT_TEXT, problems,
                                 tariff=student.tariff or ""))
            continue

        if has_left(student, now):
            left.append(student)
            continue

        opens = payments.window_opens_at(student, period)
        cutoff = payments.window_cutoff(student, period)
        if now < opens:
            state, detail = STATE_UPCOMING, f"Окно откроется {_msk_day(opens)}"
        elif now < cutoff:
            state, detail = STATE_OPEN, f"Окно до {payments.window_last_day(student, period).strftime('%d.%m')}"
        else:
            state, detail = STATE_OVERDUE, f"Окно закрылось {_msk_day(cutoff)}"
        rows.append(BoardRow(student, state, window, payments.price_kop(db, student), None, detail, problems,
                             tariff=student.tariff or ""))

    rows.sort(key=lambda r: STATE_ORDER.index(r.state))
    counts = {state: sum(1 for r in rows if r.state == state) for state in STATE_ORDER}
    return Board(
        period, rows, no_window, counts, paid_total,
        paid_count=sum(1 for r in rows if r.state == STATE_PAID and not r.no_payment),
        no_payment_count=sum(1 for r in rows if r.no_payment),
        debt_total_kop=sum(r.amount_kop or 0 for r in rows if r.state == STATE_OVERDUE),
        by_tariff=_tariff_totals(rows),
        left=left,
    )


def filter_rows(rows: list[BoardRow], *, status: str = "", q: str = "") -> list[BoardRow]:
    """Строки таблицы под фильтр и поиск. Итоги месяца считаются по всем
    строкам — фильтр сужает только список. Поиск — общее правило
    (`services/student_search.py`): имя, фамилия, ник, номер."""
    result = []
    for row in rows:
        if status == FILTER_PAID and row.state != STATE_PAID:
            continue
        if status == FILTER_UNPAID and row.state == STATE_PAID:
            continue
        if not student_matches(row.student, q):
            continue
        result.append(row)
    return result


def _tariff_totals(rows: list[BoardRow]) -> list[TariffTotal]:
    totals: dict[str, TariffTotal] = {}
    for row in rows:
        total = totals.setdefault(row.tariff, TariffTotal(row.tariff))
        if row.no_payment:
            total.no_payment_count += 1
        elif row.state == STATE_PAID:
            total.paid_count += 1
            total.paid_kop += row.amount_kop or 0
        elif row.state == STATE_OVERDUE:
            total.debt_count += 1
            total.debt_kop += row.amount_kop or 0
        else:
            total.waiting_count += 1
    return sorted(totals.values(), key=lambda t: (-t.paid_kop, -t.paid_count, tariff_text(t.tariff)))


def months_summary(db: DBSession, now: datetime) -> list[MonthTotal]:
    """Сводка «месяц → оплатили / без платежа / должники», свежий месяц
    первым: от первого месяца с платежом до текущего (или до последнего
    оплаченного вперёд). Каждый месяц считает тот же `month_board`, что и экран.

    Месяцы до первого платежа не показываем: сентябрь 2026 закрыт технической
    отметкой запуска, денег за него нет. Долг за прошлые месяцы — по
    сегодняшним ценам и сегодняшнему составу учеников.
    """
    current = payments.month_start(now.astimezone(MSK_TZ).date())
    first, last = (
        db.query(func.min(Payment.period), func.max(Payment.period))
        .filter(Payment.kind == KIND_MONTH)
        .one()
    )
    start = min(first, current) if first else current
    end = max(last, current) if last else current
    students = accounted_students(db)
    result: list[MonthTotal] = []
    period = start
    while period <= end:
        board = month_board(db, period, now, students)
        result.append(MonthTotal(
            period=period,
            paid_count=board.paid_count,
            paid_kop=board.paid_total_kop,
            no_payment_count=board.no_payment_count,
            debt_count=board.counts[STATE_OVERDUE],
            debt_kop=board.debt_total_kop,
            waiting_count=board.counts[STATE_OPEN] + board.counts[STATE_UPCOMING],
        ))
        period = payments.add_months(period, 1)
    result.reverse()
    return result


def stats_card(db: DBSession, now: datetime) -> dict:
    """Карточка «Оплаты» в «Статистике активности» (владелец 14.10.2026:
    «перенести оплату в статистику»): итоги текущего месяца и сводка по
    месяцам — те же `months_summary`, что на экране «Оплаты». Поимённый
    список, фильтр и кнопки остаются на экране, карточка ведёт туда."""
    current = parse_month("", now)
    summary = months_summary(db, now)

    def money(count: int, amount_kop: int) -> str:
        return f"{count} · {payments.rub_text(amount_kop)}" if amount_kop else str(count)

    months = [
        {
            "value": m.period.strftime("%Y-%m"),
            "title": payments.period_label(m.period).capitalize(),
            "paid": money(m.paid_count, m.paid_kop),
            "no_payment": m.no_payment_count,
            "debt": money(m.debt_count, m.debt_kop),
            "waiting": m.waiting_count,
            "is_current": m.period == current,
        }
        for m in summary
    ]
    month = next((m for m in summary if m.period == current), None)
    return {
        "month_value": current.strftime("%Y-%m"),
        "month_title": payments.period_label(current),
        "paid_count": month.paid_count if month else 0,
        "paid_text": payments.rub_text(month.paid_kop) if month and month.paid_kop else "",
        "debt_count": month.debt_count if month else 0,
        "debt_text": payments.rub_text(month.debt_kop) if month and month.debt_kop else "",
        "waiting_count": month.waiting_count if month else 0,
        "no_payment_count": month.no_payment_count if month else 0,
        "months": months,
    }


# ── Excel ────────────────────────────────────────────────────────────────────

def _comment(row: BoardRow) -> str:
    parts = [row.payment.note] if row.payment is not None and row.payment.note else []
    return "; ".join(parts + row.problems)


def _rub(amount_kop: int | None) -> float | None:
    """Рубли для ячейки: в итогах ноль остаётся нулём, пусто — только где суммы нет."""
    return amount_kop / 100 if amount_kop is not None else None


def build_month_xlsx(board: Board, summary: list[MonthTotal]) -> bytes:
    """Выгрузка экрана «Оплаты» за месяц (владелец 13.10.2026): ученики с
    контактами и подробностями оплаты, итоги по тарифам, должники, сводка
    по месяцам. Оформление — как у `staff_dashboard.build_tariff_registration_xlsx`.

    Почта и Telegram здесь есть по решению владельца 13.10.2026, хотя
    выгрузка «Регистрации по тарифам» их нарочно не несёт: файл видят
    только суперадмин и ГП (`payments.is_payments_staff`)."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    def sheet(title: str, columns: tuple[tuple[str, int], ...], rows: list[list], first: bool = False) -> None:
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append([name for name, _ in columns])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="2563EB")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for index, (_, width) in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions

    def student_cells(student: User) -> list:
        tg = normalize_tg_username(student.tg_username or "")
        return [display_name(student), student.email or "", f"@{tg}" if tg else ""]

    sheet("Ученики", (
        ("Ученик", 30), ("Почта", 28), ("Telegram", 20), ("Тариф", 22), ("Статус", 20),
        ("Сумма, ₽", 12), ("Дата оплаты", 13), ("Способ", 24), ("Заказ Продамуса", 16),
        ("Окно", 8), ("Комментарий", 50), ("ID", 7),
    ), [
        student_cells(row.student) + [
            tariff_text(row.tariff),
            NO_PAYMENT_TEXT if row.no_payment else STATE_LABELS[row.state],
            _rub(row.amount_kop),
            row.paid_on,
            row.method,
            (row.payment.prodamus_order_id or "") if row.payment else "",
            row.window_text,
            _comment(row),
            row.student.id,
        ]
        for row in board.rows
    ] + [
        # Ушедшие — в конце листа, без суммы: в итогах их нет (`has_left`).
        student_cells(s) + [
            tariff_text(s.tariff), LEFT_TEXT, None, "", "", "",
            f"{s.pay_window_start}–{s.pay_window_end}",
            f"Доступ до {msk_date_text(s.access_until)}", s.id,
        ]
        for s in board.left
    ], first=True)

    sheet("По тарифам", (
        ("Тариф", 24), ("Оплатили", 11), ("Оплачено, ₽", 14), ("Без платежа", 13),
        ("Должники", 11), ("Долг, ₽", 12), ("Ждём оплату", 13),
    ), [
        [tariff_text(t.tariff), t.paid_count, _rub(t.paid_kop), t.no_payment_count,
         t.debt_count, _rub(t.debt_kop), t.waiting_count]
        for t in board.by_tariff
    ] + [[
        "Всего", board.paid_count, _rub(board.paid_total_kop), board.no_payment_count,
        board.counts[STATE_OVERDUE], _rub(board.debt_total_kop),
        board.counts[STATE_OPEN] + board.counts[STATE_UPCOMING],
    ]])

    sheet("Должники", (
        ("Ученик", 30), ("Почта", 28), ("Telegram", 20), ("Тариф", 22), ("Долг, ₽", 12),
        ("Окно", 8), ("Комментарий", 40), ("ID", 7),
    ), [
        student_cells(row.student) + [tariff_text(row.tariff), _rub(row.amount_kop), row.window_text,
                              _comment(row), row.student.id]
        for row in board.rows if row.state == STATE_OVERDUE
    ])

    sheet("По месяцам", (
        ("Месяц", 16), ("Оплатили", 11), ("Оплачено, ₽", 14), ("Без платежа", 13),
        ("Должники", 11), ("Долг, ₽", 12), ("Ждём оплату", 13),
    ), [
        [payments.period_label(m.period).capitalize(), m.paid_count, _rub(m.paid_kop),
         m.no_payment_count, m.debt_count, _rub(m.debt_kop), m.waiting_count]
        for m in summary
    ])

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
